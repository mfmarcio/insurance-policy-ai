from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
DIAGRAM = DOCS / 'architecture.png'
OUT = DOCS / 'Relatorio_Tecnico_PolicyCompare_AI.docx'


def font(size=28, bold=False):
    path = '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf' if bold else '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    return ImageFont.truetype(path, size)


def rounded(draw, box, text, fill, outline=(80,80,80), fnt=None):
    draw.rounded_rectangle(box, radius=18, fill=fill, outline=outline, width=2)
    x1,y1,x2,y2=box
    bbox=draw.multiline_textbbox((0,0),text,font=fnt,spacing=4,align='center')
    tw=bbox[2]-bbox[0]; th=bbox[3]-bbox[1]
    draw.multiline_text(((x1+x2-tw)/2,(y1+y2-th)/2),text,font=fnt,fill=(25,25,25),spacing=4,align='center')


def arrow(draw, x1,y1,x2,y2):
    draw.line((x1,y1,x2,y2), fill=(60,60,60), width=5)
    import math
    ang=math.atan2(y2-y1,x2-x1)
    l=16
    pts=[(x2,y2),(x2-l*math.cos(ang-0.5),y2-l*math.sin(ang-0.5)),(x2-l*math.cos(ang+0.5),y2-l*math.sin(ang+0.5))]
    draw.polygon(pts, fill=(60,60,60))


def build_diagram():
    img=Image.new('RGB',(1500,1050),'white')
    d=ImageDraw.Draw(img)
    title=font(40,True); boxf=font(24,True); small=font(19)
    d.text((750,35),'PolicyCompare AI - Arquitetura do MVP',font=title,anchor='ma',fill=(20,20,20))
    boxes=[
        ((525,110,975,205),'Interface Web\nStreamlit',(227,242,253)),
        ((525,250,975,345),'Intake Agent\nValidação + hash do documento',(235,246,235)),
        ((525,390,975,485),'Extraction Agent\nPyMuPDF + OCR Tesseract',(255,243,224)),
        ((525,530,975,625),'Clause Agent\nRetrieval TF-IDF + Ollama',(250,235,245)),
        ((525,670,975,765),'Structuring Agent\nPydantic + schema canônico',(238,238,250)),
        ((50,825,400,930),'SQLite\nApólices estruturadas',(245,245,245)),
        ((575,825,925,930),'Query Agent\nPerguntas + evidências',(235,246,235)),
        ((1100,825,1450,930),'Comparison + Report Agents\nRegras + LLM + PDF',(245,245,245)),
    ]
    for box,text,fill in boxes: rounded(d,box,text,fill,fnt=boxf)
    for a,b in [((750,205),(750,250)),((750,345),(750,390)),((750,485),(750,530)),((750,625),(750,670)),((650,765),(225,825)),((225,930),(700,930)),((400,875),(575,875)),((400,900),(1100,900))]: arrow(d,*a,*b)
    d.text((225,955),'Persistência',font=small,anchor='ma',fill=(70,70,70))
    d.text((750,955),'Consulta rastreável',font=small,anchor='ma',fill=(70,70,70))
    d.text((1275,955),'Comparação e relatório',font=small,anchor='ma',fill=(70,70,70))
    img.save(DIAGRAM)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('Página ')
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)


def add_heading(doc, text, level=1):
    p=doc.add_heading(text, level=level)
    p.paragraph_format.space_before=Pt(10)
    p.paragraph_format.space_after=Pt(5)
    return p


def add_bullets(doc, items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet')
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table=doc.add_table(rows=1, cols=len(headers))
    table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=table.rows[0]
    set_repeat_table_header(hdr)
    for i,h in enumerate(headers):
        cell=hdr.cells[i]; cell.text=h; shade_cell(cell,'D9EAF7'); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in cell.paragraphs[0].runs: r.bold=True
    for row in rows:
        cells=table.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v)
            cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    return table


build_diagram()
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(0.7); sec.bottom_margin=Inches(0.7); sec.left_margin=Inches(0.75); sec.right_margin=Inches(0.75)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5)
styles['Title'].font.name='Aptos Display'; styles['Title'].font.size=Pt(28); styles['Title'].font.bold=True
for s in ['Heading 1','Heading 2','Heading 3']:
    styles[s].font.name='Aptos Display'
styles['Heading 1'].font.color.rgb=RGBColor(31,78,121)
styles['Heading 2'].font.color.rgb=RGBColor(47,84,150)

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(75)
r=p.add_run('POLICYCOMPARE AI'); r.bold=True; r.font.size=Pt(30); r.font.color.rgb=RGBColor(31,78,121)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Plataforma Inteligente para Análise e Comparação de Apólices'); r.font.size=Pt(18)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(20)
r=p.add_run('Relatório Técnico - MVP'); r.bold=True; r.font.size=Pt(15)
doc.add_picture(str(DIAGRAM), width=Inches(6.5))
doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(25)
p.add_run('Versão de demonstração - Outubro de 2026').italic=True
footer=sec.footer.paragraphs[0]; add_page_number(footer)
doc.add_page_break()

add_heading(doc,'1. Resumo executivo',1)
doc.add_paragraph(
    'O PolicyCompare AI é um MVP local para automatizar etapas repetitivas da leitura e comparação de apólices de seguro. '
    'A solução recebe documentos em PDF ou imagem, extrai o conteúdo, identifica informações relevantes por meio de IA Generativa, '
    'estrutura os resultados em um modelo canônico, persiste os dados e compara duas apólices apresentando diferenças e evidências de origem.'
)
doc.add_paragraph(
    'A arquitetura foi desenhada como um fluxo multiagente. Cada componente possui uma responsabilidade explícita, reduzindo acoplamento e permitindo '
    'substituir OCR, modelo de linguagem, banco ou mecanismo de recuperação sem reescrever todo o sistema. O princípio adotado é: IA interpreta; código '
    'valida e compara determinísticamente sempre que possível.'
)
add_bullets(doc,[
    'Leitura de PDF e imagens, com OCR gratuito como fallback.',
    'Uso de LLM local via Ollama, evitando dependência obrigatória de APIs pagas.',
    'Saída estruturada validada por JSON Schema/Pydantic.',
    'Rastreabilidade por página, evidência textual e score de confiança.',
    'Comparação A x B com classificação das diferenças.',
    'Interface Streamlit e geração de relatório comparativo em PDF.',
])

add_heading(doc,'2. Problema e objetivo',1)
doc.add_paragraph(
    'Apólices são documentos extensos, heterogêneos e frequentemente redigidos em linguagem jurídica. A comparação manual exige localizar coberturas, '
    'exclusões, franquias, limites de indenização, vigência e cláusulas especiais em diferentes pontos do documento. O objetivo do projeto é reduzir o '
    'tempo gasto nessa leitura, sem eliminar a necessidade de validação de especialistas.'
)
add_heading(doc,'2.1 Escopo do MVP',2)
add_bullets(doc,[
    'Receber múltiplos arquivos PDF, PNG, JPG, JPEG, TIFF/TIF.',
    'Extrair texto nativo de PDF e aplicar OCR quando o texto for ausente ou insuficiente.',
    'Identificar dados gerais, coberturas, exclusões e cláusulas especiais.',
    'Persistir o resultado estruturado em SQLite.',
    'Comparar duas apólices e destacar equivalências, diferenças e itens exclusivos.',
    'Exibir evidências e páginas de origem.',
    'Gerar um PDF comparativo para demonstração.',
])

add_heading(doc,'3. Arquitetura da solução',1)
doc.add_picture(str(DIAGRAM), width=Inches(6.8)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('Figura 1 - Arquitetura lógica do MVP.').alignment=WD_ALIGN_PARAGRAPH.CENTER
add_heading(doc,'3.1 Separação de responsabilidades',2)
add_table(doc,['Componente','Responsabilidade','Entrada','Saída'],[
    ('Intake Agent','Validação de formato, arquivo e identificação por hash.','PDF/Imagem','document_id'),
    ('Extraction Agent','Extração de texto e OCR de páginas digitalizadas.','Documento','Texto por página'),
    ('Chunking/Recovery','Divisão do texto e seleção de trechos relevantes.','Texto','Chunks ranqueados'),
    ('Clause Agent','Interpretação semântica por LLM local.','Chunks','Dados de apólice'),
    ('Structuring Agent','Validação e normalização do schema.','Dados do LLM','PolicyData'),
    ('Repository','Persistência estruturada.','PolicyData','SQLite'),
    ('Query Agent','Perguntas sobre uma apólice estruturada com evidências.','Pergunta + PolicyData','Resposta + evidências'),
    ('Comparison Agent','Comparação objetiva e semântica.','Apólice A e B','PolicyComparison'),
    ('Report Agent','Relatório comparativo.','Comparação','PDF'),
])

add_heading(doc,'4. Tecnologias utilizadas',1)
add_table(doc,['Tecnologia','Uso','Justificativa'],[
    ('Python 3.11+','Back-end e agentes','Ecossistema maduro para IA, OCR, documentos e web.'),
    ('Streamlit','Interface de demonstração','Permite entregar interface funcional com pouco código.'),
    ('Ollama','Execução de LLM local','Privacidade, ausência de chave externa e troca simples de modelo.'),
    ('Qwen local (configurável)','Extração semântica e resumo','Modelo local adequado ao protótipo; pode ser substituído por outro disponível no Ollama.'),
    ('PyMuPDF','Leitura de PDF','Extração textual rápida e renderização de páginas.'),
    ('Tesseract OCR','OCR gratuito','Fallback para imagens e PDFs digitalizados.'),
    ('Pydantic','Schemas e validação','Contrato forte entre LLM e aplicação.'),
    ('LangGraph','Orquestração','Workflow explícito com nós/agentes especializados.'),
    ('SQLite','Persistência','Zero infraestrutura adicional para o MVP.'),
    ('scikit-learn / TF-IDF','Recuperação local','Reduz o contexto enviado ao LLM sem exigir vector DB.'),
    ('ReportLab','PDF comparativo','Geração programática do relatório final.'),
])

add_heading(doc,'5. Descrição dos agentes',1)
for title, body, bullets in [
    ('5.1 Intake Agent','Responsável pela porta de entrada do fluxo. Valida extensão, existência e tamanho do arquivo e produz identificador estável derivado de SHA-256.',[
        'Evita processar formatos não suportados.','Permite reconhecer o mesmo documento em execuções diferentes.']),
    ('5.2 Extraction Agent','Tenta primeiro a extração textual nativa. Se uma página possuir pouco texto, renderiza a página e aplica OCR.',[
        'PDF textual: PyMuPDF.','Imagem/PDF digitalizado: Tesseract.','Resultado preserva número da página e método de extração.']),
    ('5.3 Clause Agent','Seleciona chunks potencialmente relevantes e solicita ao LLM local uma resposta estruturada.',[
        'Não deve inventar campos ausentes.','Cada cobertura/exclusão deve carregar evidência textual e página.','A saída é validada contra schema Pydantic.']),
    ('5.4 Structuring Agent','Converte a extração semântica para o modelo canônico PolicyData.',[
        'Desacopla o schema persistido do formato bruto do LLM.','Centraliza defaults, validações e tipos.']),
    ('5.5 Query Agent','Responde perguntas em linguagem natural sobre os dados já extraídos de uma apólice.',[
        'Usa exclusivamente o PolicyData como fonte.','Retorna evidências existentes e sinaliza quando a resposta é inconclusiva.']),
    ('5.6 Comparison Agent','Combina pareamento textual com regras determinísticas.',[
        'Limite maior pode ser classificado como mais favorável.','Franquia menor pode ser classificada como mais favorável.','Itens presentes somente em uma apólice são destacados.','Exclusões próximas são pareadas por similaridade.']),
    ('5.7 Report Agent','Produz um documento comparativo contendo resumo, tabela de diferenças e aviso de validação humana.',[
        'Entrega um artefato compartilhável.','Evita que o resultado fique restrito à interface web.']),
]:
    add_heading(doc,title,2); doc.add_paragraph(body); add_bullets(doc,bullets)

add_heading(doc,'6. Fluxo completo de processamento',1)
steps=[
    'Usuário envia uma ou mais apólices pela interface.',
    'Intake Agent valida os arquivos e gera document_id.',
    'Extraction Agent extrai texto por página; OCR é acionado quando necessário.',
    'Chunking Service fragmenta o texto preservando página.',
    'Retrieval Service usa TF-IDF para selecionar trechos relacionados a coberturas, exclusões, franquias, limites e dados gerais.',
    'Clause Agent envia somente os trechos selecionados ao Ollama e solicita JSON estruturado.',
    'Pydantic valida a resposta; o Structuring Agent cria o PolicyData.',
    'PolicyRepository persiste o JSON no SQLite.',
    'Usuário pode consultar uma apólice estruturada em linguagem natural por meio do Query Agent.',
    'Usuário seleciona duas apólices na aba de comparação.',
    'Comparison Agent pareia elementos e classifica diferenças.',
    'O LLM produz um resumo executivo, com fallback determinístico se indisponível.',
    'Report Agent gera o PDF comparativo e a interface apresenta evidências por página.',
]
for n,s in enumerate(steps,1):
    p=doc.add_paragraph(style='List Number'); p.add_run(s)

add_heading(doc,'7. Modelo de dados e rastreabilidade',1)
doc.add_paragraph(
    'O schema canônico foi desenhado para não armazenar apenas a conclusão do modelo. Informações relevantes carregam Evidence com página, texto e confiança. '
    'Esse desenho permite auditoria, revisão humana e futuras métricas de qualidade.'
)
add_table(doc,['Entidade','Campos principais'],[
    ('PolicyData','policy_id, document_name, insurer, policy_number, insured_name, product, vigência'),
    ('Coverage','name, description, limit_value, currency, deductible_value, deductible_text, evidence'),
    ('Exclusion','category, description, evidence'),
    ('SpecialClause','category, title, description, evidence'),
    ('Evidence','page, text, confidence'),
    ('PolicyComparison','policy_a_id, policy_b_id, summary, items'),
    ('ComparisonItem','dimension, item, A, B, status, rationale, evidence_a, evidence_b'),
])

add_heading(doc,'8. Estratégia de comparação',1)
doc.add_paragraph('A comparação não é delegada integralmente ao LLM. O MVP prioriza regras explícitas para reduzir variabilidade e facilitar testes.')
add_table(doc,['Classificação','Significado'],[
    ('EQUIVALENTE','Condições equivalentes nos campos comparados.'),
    ('MAIS_FAVORAVEL_A','Critério quantitativo favorece A, por exemplo maior limite ou menor franquia.'),
    ('MAIS_FAVORAVEL_B','Critério quantitativo favorece B.'),
    ('SOMENTE_A','Item identificado somente na apólice A.'),
    ('SOMENTE_B','Item identificado somente na apólice B.'),
    ('DIFERENTE','Conteúdo presente em ambas, mas materialmente diferente.'),
    ('INCONCLUSIVO','Dados insuficientes para uma conclusão segura.'),
])

add_heading(doc,'9. Tratamento de erros, segurança e configuração',1)
add_bullets(doc,[
    'Arquivo inválido ou vazio interrompe o fluxo antes do processamento de IA.',
    'Falhas de conexão ou timeout do Ollama são convertidas em erro controlado.',
    'A resposta estruturada do LLM é validada por Pydantic; JSON incompatível não é persistido silenciosamente.',
    'OCR é fallback e não é aplicado indiscriminadamente a todo PDF.',
    'Configurações ficam em .env; o arquivo real é ignorado pelo Git.',
    'O MVP não exige chaves de API externas e pode operar integralmente no computador do avaliador.',
    'Documentos reais podem conter dados sensíveis; para produção seriam necessários autenticação, controle de acesso, criptografia, retenção e trilha de auditoria.',
])

add_heading(doc,'10. Instalação e execução',1)
add_bullets(doc,[
    'Instalar Python 3.11 ou 3.12.',
    'Criar ambiente virtual e executar pip install -r requirements.txt.',
    'Instalar Ollama e executar: ollama pull qwen3:4b.',
    'Instalar Tesseract OCR e o idioma português, garantindo tesseract.exe no PATH.',
    'Copiar .env.example para .env.',
    'Executar python run.py ou streamlit run frontend/streamlit_app.py.',
    'Abrir http://localhost:8501.',
])
add_heading(doc,'10.1 Roteiro de demonstração',2)
add_bullets(doc,[
    'Criar os dois PDFs sintéticos pelo botão da barra lateral.',
    'Enviar os dois arquivos pela aba Processar.',
    'Acompanhar a estrutura JSON extraída.',
    'Abrir a aba Apólices e conferir coberturas/exclusões.',
    'Usar a aba Consultar para fazer uma pergunta e revisar evidências.',
    'Selecionar A e B na aba Comparar.',
    'Conferir classificação, evidências e páginas.',
    'Baixar o PDF comparativo.',
])

add_heading(doc,'11. Testes e validação',1)
doc.add_paragraph('O pacote inclui testes automatizados para pontos determinísticos e foi submetido a smoke tests locais durante a geração deste entregável.')
add_table(doc,['Teste','Resultado esperado'],[
    ('Intake ID','Mesmo arquivo produz o mesmo identificador.'),
    ('Comparação de limite','Maior limite em A resulta em MAIS_FAVORAVEL_A.'),
    ('PDF sintético','PyMuPDF extrai texto dos PDFs de demonstração.'),
    ('Relatório comparativo','ReportLab gera PDF não vazio com a tabela de comparação.'),
])

doc.add_paragraph('Comando: pytest -q')

add_heading(doc,'12. Limitações conhecidas',1)
add_bullets(doc,[
    'A ontologia é genérica; termos de diferentes ramos de seguro podem exigir especialização.',
    'Tabelas muito complexas, documentos com múltiplas colunas e baixa qualidade de digitalização podem reduzir a extração.',
    'TF-IDF não possui a mesma capacidade semântica de embeddings modernos.',
    'Modelos locais pequenos podem perder nuances jurídicas em cláusulas longas.',
    'A heurística de pareamento de coberturas usa similaridade textual e pode necessitar de taxonomia controlada.',
    'Não existem autenticação, multiusuário, RBAC, criptografia ou alta disponibilidade no MVP.',
    'O resultado deve ser validado por profissional qualificado e não constitui aconselhamento jurídico.',
])

add_heading(doc,'13. Evolução futura',1)
add_bullets(doc,[
    'Substituir TF-IDF por embeddings multilíngues e Qdrant/pgvector.',
    'Adicionar taxonomia de coberturas e exclusões por ramo de seguro.',
    'Criar etapa human-in-the-loop para itens abaixo de um limiar de confiança.',
    'Adicionar OCR/layout especializado para tabelas, formulários e documentos escaneados.',
    'Executar extrações por página/seção em paralelo.',
    'Criar benchmark com apólices rotuladas e métricas precision, recall e F1 por campo.',
    'Versionar prompts, modelos e resultados para auditoria.',
    'Adicionar FastAPI, autenticação, trilha de auditoria e storage seguro.',
    'Adicionar exportação JSON/Excel e comparação de três ou mais apólices.',
    'Adicionar perguntas em linguagem natural com RAG e evidência clicável.',
])

add_heading(doc,'14. Estrutura do repositório',1)
code='''insurance-policy-ai/\n├── app/\n│   ├── agents/\n│   ├── config/\n│   ├── models/\n│   ├── repositories/\n│   ├── services/\n│   └── workflows/\n├── frontend/\n├── tests/\n├── data/\n├── samples/\n├── docs/\n├── requirements.txt\n├── docker-compose.yml\n├── .env.example\n└── README.md'''
p=doc.add_paragraph(); r=p.add_run(code); r.font.name='Consolas'; r.font.size=Pt(9)

add_heading(doc,'15. Conclusão',1)
doc.add_paragraph(
    'O MVP atende aos requisitos centrais do desafio: recebe PDFs e imagens, extrai conteúdo, utiliza IA Generativa local, estrutura dados, persiste resultados, '
    'compara pelo menos duas apólices e apresenta diferenças em uma interface demonstrável. A arquitetura modular e multiagente deixa explícito onde cada '
    'responsabilidade reside e oferece um caminho direto para incorporar embeddings, OCR avançado, governança e integrações sem reconstruir a solução.'
)

add_heading(doc,'Apêndice A - Critérios de aceitação do MVP',1)
add_table(doc,['Requisito','Atendimento'],[
    ('Leitura PDF/imagem','Sim - PyMuPDF + Tesseract'),
    ('Extração automática','Sim - OCR + LLM'),
    ('Dados estruturados','Sim - Pydantic/JSON'),
    ('Comparação de 2 apólices','Sim'),
    ('Diferenças ao usuário','Sim - tabela, status e evidências'),
    ('IA Generativa','Sim - Ollama'),
    ('Interface de demonstração','Sim - Streamlit'),
    ('Agentes especializados','Sim - 7 agentes + workflow'),
    ('Tratamento de erros','Sim - validações e exceções controladas'),
    ('Credenciais ocultas','Sim - .env/.gitignore; por padrão sem API externa'),
    ('Instruções de instalação','Sim - README'),
])

doc.save(OUT)
print(OUT)
